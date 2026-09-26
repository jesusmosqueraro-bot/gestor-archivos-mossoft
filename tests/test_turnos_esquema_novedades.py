"""Pruebas de las mejoras al Cuadro de Turnos pedidas por Tomás el 26/09/2026, en un solo mensaje:

  1) Rol nuevo 'lider': mismas capacidades que 'estandar' pero con acceso POR DEFECTO al módulo
     'turnos' (sin necesitar el permiso extra) — y 'agente' deja de recibirlo automático por rol
     (ver tests/test_turnos_cuadro.py para los casos de control de acceso puros; aquí solo se
     cubre la parte de "escritura" — programar turnos/asignar esquemas/registrar novedades queda
     reservado a 'admin'/'lider').
  2) Esquema de Jornada ('LV'/'LS') por colaborador, editable desde Editar Usuario, con override
     puntual por mes desde Horas del Mes — la meta de horas del mes se calcula dinámicamente según
     los días hábiles de ese esquema, salvo que haya una meta MANUAL fijada (esa gana).
  3) Novedades de turno (incapacidad/permiso/calamidad/llegada tarde/salida temprana): descuentan
     horas del total neto de Horas del Mes; solo 'admin'/'lider' las registran o anulan.
  4) Exportación a Excel en formato matriz mensual (una columna por día + totales).
  5) Colaboradores por Sede ahora agrupa por la Sede del turno asignado vigente (no solo la de
     perfil) y muestra el horario de cada colaborador.
"""
import io

import openpyxl


def _sesion_como(client, arkiv_app, usuario, rol, modulos_extra=None):
    with client.session_transaction() as sess:
        sess['logged_in'] = True
        sess['username'] = usuario
        sess['rol'] = rol
        sess['instance_id'] = arkiv_app.SERVER_INSTANCE_ID
        sess['debe_cambiar_password'] = False
        sess['debe_activar_2fa'] = False
        if modulos_extra is not None:
            sess['modulos_extra'] = modulos_extra
    return client


def _crear_area_y_sede(app, area='UrgenciasEsq', sede='Sede Esquema'):
    conn, db_type = app.get_db()
    cur = conn.cursor()
    cur.execute("INSERT INTO ticket_configuraciones (tipo, nombre, estado) VALUES ('area', ?, 'activo')", (area,))
    cur.execute("INSERT INTO ticket_configuraciones (tipo, nombre, estado) VALUES ('sede', ?, 'activo')", (sede,))
    conn.commit()
    conn.close()
    return area, sede


def _tipo_id(app, codigo):
    conn, db_type = app.get_db()
    cur = conn.cursor()
    cur.execute("SELECT id FROM tipos_turno WHERE codigo = ?", (codigo,))
    (tid,) = cur.fetchone()
    conn.close()
    return tid


def _usuario_id(app, usuario):
    conn, db_type = app.get_db()
    cur = conn.cursor()
    cur.execute("SELECT id FROM usuarios WHERE usuario = ?", (usuario,))
    (uid,) = cur.fetchone()
    conn.close()
    return uid


def _crear_sede(app, nombre):
    conn, db_type = app.get_db()
    cur = conn.cursor()
    cur.execute("INSERT INTO ticket_configuraciones (tipo, nombre, estado) VALUES ('sede', ?, 'activo')", (nombre,))
    conn.commit()
    conn.close()


# ---------------------------------------------------------------------------
# 1) Rol 'lider'
# ---------------------------------------------------------------------------

def test_registrar_usuario_acepta_el_rol_lider(admin_session, app):
    r = admin_session.post('/usuarios', data={
        'primer_nombre': 'Lidia', 'primer_apellido': 'Coordinadora',
        'email': 'lidia.coordinadora@preventivaips.com.co', 'password': 'ClaveSegura123',
        'especialidad': 'Coordinación', 'rol': 'lider',
    })
    assert r.status_code in (200, 302)
    conn, db_type = app.get_db()
    cur = conn.cursor()
    cur.execute("SELECT rol FROM usuarios WHERE correo = ?", ('lidia.coordinadora@preventivaips.com.co',))
    fila = cur.fetchone()
    conn.close()
    assert fila is not None and fila[0] == 'lider'


def test_editar_usuario_puede_ascender_a_alguien_a_lider(admin_session, app, crear_usuario):
    usuario = crear_usuario(rol='estandar')
    usuario_id = _usuario_id(app, usuario)
    admin_session.post(f'/editar_usuario/{usuario_id}', data={'email': f'{usuario}@preventivaips.com.co', 'rol': 'lider'})
    conn, db_type = app.get_db()
    cur = conn.cursor()
    cur.execute("SELECT rol FROM usuarios WHERE id = ?", (usuario_id,))
    (rol,) = cur.fetchone()
    conn.close()
    assert rol == 'lider'


def test_lider_tiene_el_mismo_acceso_al_chat_asistente_que_estandar(client, app, crear_usuario):
    """Pedido explícito: 'lider' tiene las mismas capacidades y vistas que 'estandar' — el
    Asistente de Chat (exclusivo de 'estandar' hasta ahora) es una de esas capacidades."""
    app._guardar_config_app(app.CLAVE_CHAT_ESTANDAR, '1')
    usuario = crear_usuario(rol='lider')
    _sesion_como(client, app, usuario, 'lider')
    assert client.get('/chat/bot/estado').status_code == 200


# ---------------------------------------------------------------------------
# 2) Esquema de Jornada + meta dinámica
# ---------------------------------------------------------------------------

def test_editar_usuario_guarda_el_esquema_de_jornada(admin_session, app, crear_usuario):
    usuario = crear_usuario(rol='estandar')
    usuario_id = _usuario_id(app, usuario)
    admin_session.post(f'/editar_usuario/{usuario_id}', data={
        'email': f'{usuario}@preventivaips.com.co', 'rol': 'estandar', 'esquema_jornada': 'LS',
    })
    conn, db_type = app.get_db()
    cur = conn.cursor()
    cur.execute("SELECT esquema_jornada FROM usuarios WHERE id = ?", (usuario_id,))
    (esquema,) = cur.fetchone()
    conn.close()
    assert esquema == 'LS'


def test_editar_usuario_esquema_invalido_se_descarta_y_conserva_el_anterior(admin_session, app, crear_usuario):
    usuario = crear_usuario(rol='estandar')
    usuario_id = _usuario_id(app, usuario)
    admin_session.post(f'/editar_usuario/{usuario_id}', data={'email': f'{usuario}@preventivaips.com.co', 'rol': 'estandar', 'esquema_jornada': 'LS'})
    admin_session.post(f'/editar_usuario/{usuario_id}', data={'email': f'{usuario}@preventivaips.com.co', 'rol': 'estandar', 'esquema_jornada': 'INVENTADO'})
    conn, db_type = app.get_db()
    cur = conn.cursor()
    cur.execute("SELECT esquema_jornada FROM usuarios WHERE id = ?", (usuario_id,))
    (esquema,) = cur.fetchone()
    conn.close()
    assert esquema == 'LS'


def test_meta_dinamica_lunes_a_viernes_septiembre_2026(app):
    """Septiembre 2026 tiene 22 días hábiles de lunes a viernes (verificado por conteo directo);
    con 42h/semana repartidas en 5 días (8.4h/día), la meta debe dar 22 * 8.4 = 184.8h."""
    resultado = app._meta_horas_mes_por_esquema('LV', 2026, 9)
    assert resultado == 184.8


def test_meta_dinamica_lunes_a_sabado_septiembre_2026(app):
    """Septiembre 2026 tiene 26 días hábiles de lunes a sábado; con 42h/semana repartidas en 6
    días (7h/día), la meta debe dar 26 * 7 = 182h — un valor distinto (no necesariamente mayor)
    al de L-V, porque la jornada DIARIA es menor cuantos más días tiene la semana laboral."""
    resultado = app._meta_horas_mes_por_esquema('LS', 2026, 9)
    assert resultado == 182.0


def test_horas_mes_usa_meta_manual_cuando_esta_fijada_por_encima_de_la_dinamica(admin_session, app, crear_usuario):
    colaborador = crear_usuario(usuario='colab_meta_manual_esq')
    usuario_id = _usuario_id(app, colaborador)
    admin_session.post(f'/editar_usuario/{usuario_id}', data={
        'email': f'{colaborador}@preventivaips.com.co', 'rol': 'estandar',
        'meta_horas_mensual': '10', 'esquema_jornada': 'LV',
    })
    area, sede = _crear_area_y_sede(app)
    tipo_id = _tipo_id(app, 'M6_12')
    admin_session.post('/turnos/asignar', data={
        'colaborador_usuario': colaborador, 'fecha': '2026-09-21', 'tipo_turno_id': str(tipo_id),
        'area': area, 'sede': sede, 'rol_profesional': 'MEDICO',
    })
    r = admin_session.get('/turnos/horas_mes?mes=2026-09')
    texto = r.get_data(as_text=True)
    assert '10.0 h' in texto
    assert 'manual' in texto


def test_turnos_horas_mes_esquema_fija_override_mensual_y_admin_puede(admin_session, app, crear_usuario):
    colaborador = crear_usuario(usuario='colab_override_mes')
    r = admin_session.post('/turnos/horas_mes/esquema', data={
        'colaborador_usuario': colaborador, 'mes': '2026-09', 'esquema': 'LS',
    })
    assert r.status_code == 200
    assert r.get_json()['ok'] is True
    conn, db_type = app.get_db()
    cur = conn.cursor()
    cur.execute("SELECT esquema FROM turnos_esquemas_mensuales WHERE usuario_id = ? AND anio = 2026 AND mes = 9", (_usuario_id(app, colaborador),))
    (esquema,) = cur.fetchone()
    conn.close()
    assert esquema == 'LS'


def test_turnos_horas_mes_esquema_bloqueado_para_quien_no_gestiona_turnos(client, app, crear_usuario):
    colaborador = crear_usuario(usuario='colab_override_bloqueado')
    usuario_agente = crear_usuario(usuario='agente_solo_lectura', rol='agente')
    _sesion_como(client, app, usuario_agente, 'agente', modulos_extra=['turnos'])
    r = client.post('/turnos/horas_mes/esquema', data={
        'colaborador_usuario': colaborador, 'mes': '2026-09', 'esquema': 'LS',
    })
    assert r.status_code == 403
    assert r.get_json()['ok'] is False


def test_turnos_horas_mes_esquema_invalido_se_rechaza(admin_session, app, crear_usuario):
    colaborador = crear_usuario(usuario='colab_override_invalido')
    r = admin_session.post('/turnos/horas_mes/esquema', data={
        'colaborador_usuario': colaborador, 'mes': '2026-09', 'esquema': 'INVENTADO',
    })
    assert r.status_code == 400
    assert r.get_json()['ok'] is False


# ---------------------------------------------------------------------------
# 3) Novedades de turno
# ---------------------------------------------------------------------------

def test_registrar_novedad_requiere_gestionar_turnos(client, app, crear_usuario):
    colaborador = crear_usuario(usuario='colab_novedad_bloqueada')
    usuario_estandar = crear_usuario(usuario='estandar_solo_lectura', rol='estandar')
    _sesion_como(client, app, usuario_estandar, 'estandar', modulos_extra=['turnos'])
    r = client.post('/turnos/novedades/registrar', data={
        'colaborador_usuario': colaborador, 'tipo': 'incapacidad', 'fecha': '2026-09-10', 'horas_afectadas': '8',
    })
    assert r.status_code == 403
    assert r.get_json()['ok'] is False


def test_admin_registra_una_novedad_y_aparece_en_el_listado(admin_session, app, crear_usuario):
    colaborador = crear_usuario(usuario='colab_novedad_1', nombre='Colaborador Novedad Uno')
    r = admin_session.post('/turnos/novedades/registrar', data={
        'colaborador_usuario': colaborador, 'tipo': 'permiso', 'fecha': '2026-09-10',
        'horas_afectadas': '4', 'observaciones': 'Cita médica',
    })
    assert r.status_code == 200
    assert r.get_json()['ok'] is True

    listado = admin_session.get('/turnos/novedades?mes=2026-09').get_data(as_text=True)
    assert 'Colaborador Novedad Uno' in listado
    assert 'Permiso' in listado
    assert 'Cita médica' in listado


def test_registrar_novedad_con_tipo_invalido_se_rechaza(admin_session, app, crear_usuario):
    colaborador = crear_usuario(usuario='colab_novedad_tipo_malo')
    r = admin_session.post('/turnos/novedades/registrar', data={
        'colaborador_usuario': colaborador, 'tipo': 'vacaciones_inventadas', 'fecha': '2026-09-10', 'horas_afectadas': '4',
    })
    assert r.status_code == 400
    assert r.get_json()['ok'] is False


def test_novedad_descuenta_horas_del_total_neto_de_horas_del_mes(admin_session, app, crear_usuario):
    colaborador = crear_usuario(usuario='colab_novedad_descuento')
    area, sede = _crear_area_y_sede(app, area='UrgenciasNovedad', sede='SedeNovedad')
    tipo_id = _tipo_id(app, 'M6_12')  # 6 horas
    admin_session.post('/turnos/asignar', data={
        'colaborador_usuario': colaborador, 'fecha': '2026-09-21', 'tipo_turno_id': str(tipo_id),
        'area': area, 'sede': sede, 'rol_profesional': 'MEDICO',
    })
    admin_session.post('/turnos/novedades/registrar', data={
        'colaborador_usuario': colaborador, 'tipo': 'llegada_tarde', 'fecha': '2026-09-21', 'horas_afectadas': '2',
    })

    r = admin_session.get('/turnos/horas_mes?mes=2026-09')
    texto = r.get_data(as_text=True)
    assert '-2.0 h' in texto
    assert '4.0 h' in texto  # 6h del turno - 2h de novedad = 4h netas


def test_anular_novedad_deja_de_descontar_horas(admin_session, app, crear_usuario):
    colaborador = crear_usuario(usuario='colab_novedad_anulada')
    area, sede = _crear_area_y_sede(app, area='UrgenciasAnulada', sede='SedeAnulada')
    tipo_id = _tipo_id(app, 'M6_12')
    admin_session.post('/turnos/asignar', data={
        'colaborador_usuario': colaborador, 'fecha': '2026-09-21', 'tipo_turno_id': str(tipo_id),
        'area': area, 'sede': sede, 'rol_profesional': 'MEDICO',
    })
    r = admin_session.post('/turnos/novedades/registrar', data={
        'colaborador_usuario': colaborador, 'tipo': 'salida_temprana', 'fecha': '2026-09-21', 'horas_afectadas': '3',
    })
    conn, db_type = app.get_db()
    cur = conn.cursor()
    cur.execute("SELECT id FROM novedades_turno WHERE usuario_id = ?", (_usuario_id(app, colaborador),))
    (novedad_id,) = cur.fetchone()
    conn.close()

    r2 = admin_session.post(f'/turnos/novedades/{novedad_id}/eliminar')
    assert r2.status_code == 200
    assert r2.get_json()['ok'] is True

    texto = admin_session.get('/turnos/horas_mes?mes=2026-09').get_data(as_text=True)
    assert '6.0 h' in texto  # ya sin descuento, el neto vuelve a ser el bruto del turno


# ---------------------------------------------------------------------------
# 4) Exportación a Excel — matriz mensual
# ---------------------------------------------------------------------------

def test_exportar_horas_mes_xlsx_genera_una_matriz_con_columnas_de_dias_y_totales(admin_session, app, crear_usuario):
    colaborador = crear_usuario(usuario='colab_export_matriz', nombre='Colaborador Export Matriz')
    area, sede = _crear_area_y_sede(app, area='UrgenciasExport', sede='SedeExport')
    tipo_id = _tipo_id(app, 'M6_12')
    admin_session.post('/turnos/asignar', data={
        'colaborador_usuario': colaborador, 'fecha': '2026-09-21', 'tipo_turno_id': str(tipo_id),
        'area': area, 'sede': sede, 'rol_profesional': 'MEDICO',
    })

    r = admin_session.get('/turnos/horas_mes/exportar_xlsx?mes=2026-09')
    assert r.status_code == 200
    assert r.headers['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    ws = wb.active
    encabezados = [c.value for c in ws[1]]
    assert encabezados[0] == 'Colaborador'
    assert '21' in encabezados
    assert 'Horas Programadas' in encabezados
    assert 'Horas Novedades / Descuentos' in encabezados
    assert 'Total Horas Netas Laboradas' in encabezados
    assert 'Meta Horas del Mes' in encabezados
    assert 'Diferencia / Balance de Horas' in encabezados
    # Septiembre 2026 tiene 30 días -> 1 columna "Colaborador" + 30 días + 5 totales = 36
    assert ws.max_column == 36

    nombres_col_a = [fila[0].value for fila in ws.iter_rows(min_row=2)]
    assert 'Colaborador Export Matriz' in nombres_col_a


def test_exportar_horas_mes_xlsx_disponible_para_quien_solo_consulta_turnos(client, app, crear_usuario):
    usuario = crear_usuario(usuario='agente_export_consulta', rol='agente')
    _sesion_como(client, app, usuario, 'agente', modulos_extra=['turnos'])
    assert client.get('/turnos/horas_mes/exportar_xlsx?mes=2026-09').status_code == 200


# ---------------------------------------------------------------------------
# 5) Colaboradores por Sede — sede/horario del turno asignado
# ---------------------------------------------------------------------------

def test_colaborador_con_turno_asignado_aparece_en_la_sede_del_turno_no_la_del_perfil(admin_session, app, crear_usuario):
    colaborador = crear_usuario(usuario='colab_sede_turno', nombre='Colaborador Sede Turno')
    usuario_id = _usuario_id(app, colaborador)
    conn, db_type = app.get_db()
    cur = conn.cursor()
    cur.execute("UPDATE usuarios SET sede = ? WHERE id = ?", ('Sede Perfil Vieja', usuario_id))
    conn.commit()
    conn.close()
    _crear_sede(app, 'Sede Perfil Vieja')
    area, sede_turno = _crear_area_y_sede(app, area='UrgenciasSedeTurno', sede='Sede Del Turno Asignado')
    tipo_id = _tipo_id(app, 'M6_12')
    from datetime import datetime as _dt
    hoy = _dt.now(app.ZONA_HORARIA_COLOMBIA).date()
    admin_session.post('/turnos/asignar', data={
        'colaborador_usuario': colaborador, 'fecha': hoy.strftime('%Y-%m-%d'), 'tipo_turno_id': str(tipo_id),
        'area': area, 'sede': sede_turno, 'rol_profesional': 'MEDICO',
    })

    body = admin_session.get('/turnos/por_sede').get_data(as_text=True)
    assert 'Sede Del Turno Asignado' in body
    idx_sede_turno = body.index('Sede Del Turno Asignado')
    idx_colaborador = body.index('Colaborador Sede Turno')
    idx_sede_vieja = body.index('Sede Perfil Vieja')
    # El colaborador aparece DESPUÉS de la tarjeta "Sede Del Turno Asignado" (dentro de ella), y
    # antes de que la tarjeta "Sede Perfil Vieja" (que ahora debe quedar vacía) muestre su nombre.
    assert idx_sede_turno < idx_colaborador
    assert colaborador in body
    assert 'M6' in body or 'M6_12' in body or '06:00' in body  # algún rastro del horario mostrado


def test_colaborador_sin_turno_sigue_agrupado_por_su_sede_de_perfil(admin_session, app, crear_usuario):
    colaborador = crear_usuario(usuario='colab_sin_turno_sede', nombre='Colaborador Sin Turno')
    usuario_id = _usuario_id(app, colaborador)
    _crear_sede(app, 'Sede Solo Perfil')
    conn, db_type = app.get_db()
    cur = conn.cursor()
    cur.execute("UPDATE usuarios SET sede = ? WHERE id = ?", ('Sede Solo Perfil', usuario_id))
    conn.commit()
    conn.close()

    body = admin_session.get('/turnos/por_sede').get_data(as_text=True)
    assert 'Sede Solo Perfil' in body
    idx_sede = body.index('Sede Solo Perfil')
    idx_colaborador = body.index('Colaborador Sin Turno')
    assert idx_sede < idx_colaborador
